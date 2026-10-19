"""Excel hisobot: ETr/ETo jamlama varaqlari, ekin turi bo'yicha varaqlar va sifat nazorati."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_TOTAL_FILL = PatternFill("solid", fgColor="DDEBF7")


def write_report(
    path: Path,
    sheets: dict[str, pd.DataFrame],
    quality: list[tuple[str, str]],
) -> Path:
    """`sheets`: varaq nomi -> jadval (oxirgi "JAMI" qatorlari qalin chiqadi). `quality`: (ko'rsatkich, qiymat) juftliklari."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=name, index=False)
            _format_table(writer.sheets[name], frame)
        qc = pd.DataFrame(quality, columns=["Ko'rsatkich", "Qiymat"])
        qc.to_excel(writer, sheet_name="Sifat", index=False)
        _format_notes(writer.sheets["Sifat"], qc)
    return path


def _number_format(header: str) -> str | None:
    if header in {"Dalalar soni", "Hisobga olingan dalalar"} or header.endswith("m³"):
        return "#,##0"
    if header == "Maydon, ga":
        return "#,##0.00"
    if header.endswith("mm"):
        return "#,##0.0"
    if header.endswith("m³/ga"):
        return "#,##0.0"
    return None


def _format_table(sheet, frame: pd.DataFrame) -> None:
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 32
    sheet.freeze_panes = "A2"
    for idx, header in enumerate(frame.columns, start=1):
        fmt = _number_format(str(header))
        letter = get_column_letter(idx)
        sheet.column_dimensions[letter].width = max(14, min(32, len(str(header)) + 4))
        if fmt:
            for row in range(2, len(frame) + 2):
                sheet.cell(row=row, column=idx).number_format = fmt
    label_columns = [i for i, h in enumerate(frame.columns, start=1) if h == "Oy"]
    for row in range(2, len(frame) + 2):
        if any(sheet.cell(row=row, column=c).value == "JAMI" for c in label_columns):
            for cell in sheet[row]:
                cell.font = Font(bold=True)
                cell.fill = _TOTAL_FILL


def _format_notes(sheet, frame: pd.DataFrame) -> None:
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEADER_FILL
    sheet.column_dimensions["A"].width = 44
    sheet.column_dimensions["B"].width = 90
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
